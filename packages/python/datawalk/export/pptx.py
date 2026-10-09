"""PowerPoint export: one slide per tour stop. Titles, findings, KPI values, tables, and text
are native PowerPoint text (selectable); charts are the images that the walls show; speaker
notes come from the tour. Needs python-pptx: pip install 'datawalk[pptx]'."""

from __future__ import annotations

import io
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, Any

from ..core.models import Dashboard, Visualization

if TYPE_CHECKING:
    from ..reports.report import Report

INTERACTIVE_NOTE = (
    "Interactive in the DataWalk report: zoom, hover values, and exact data are not in this slide."
)
MAX_TABLE = (14, 8)  # rows, columns on one slide


def _open_template(template: Path | None):
    from pptx import Presentation
    from pptx.util import Inches

    if template is None:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
        return prs, True
    path = Path(template)
    if path.suffix.lower() == ".potx":
        # python-pptx opens presentations only: copy the template with the presentation type.
        tmp = Path(tempfile.mkdtemp()) / "template.pptx"
        with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = src.read(item)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"presentationml.template.main+xml", b"presentationml.presentation.main+xml"
                    )
                dst.writestr(item, data)
        path = tmp
    prs = Presentation(str(path))
    # Start from the template's layouts and theme, without its sample slides.
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)
    return prs, False


def _layout(prs, *names: str):
    for name in names:
        for layout in prs.slide_layouts:
            if layout.name.lower() == name.lower():
                return layout
    return prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]


class _Deck:
    def __init__(
        self,
        report: Report,
        template: Path | None,
        images: dict[str, bytes],
        levels: dict[str, str],
        watermark: str,
    ):
        from pptx.util import Inches, Pt

        self.Inches, self.Pt = Inches, Pt
        self.report = report
        self.images = images
        self.levels = levels  # from the export review: "preview" charts carry no values
        self.watermark = watermark
        self.prs, self.neutral = _open_template(template)
        self.W = self.prs.slide_width
        self.H = self.prs.slide_height
        self.margin = Inches(0.6)

    # -- building blocks ------------------------------------------------------------

    def slide(self, title: str, notes: str = ""):
        layout = (
            _layout(self.prs, "Blank")
            if self.neutral
            else _layout(self.prs, "Title Only", "Title and Content")
        )
        s = self.prs.slides.add_slide(layout)
        if s.shapes.title is not None:
            s.shapes.title.text = title
            for ph in list(s.placeholders):
                if ph.placeholder_format.idx != 0:  # empty body placeholders of the layout
                    ph._element.getparent().remove(ph._element)
            self.top = s.shapes.title.top + s.shapes.title.height + self.Inches(0.15)
        else:
            self.text(
                s,
                title,
                self.margin,
                self.Inches(0.35),
                self.W - 2 * self.margin,
                self.Inches(0.9),
                30,
                bold=True,
            )
            self.top = self.Inches(1.35)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        if self.watermark:
            x, w = self.W - self.Inches(4.6), self.Inches(4.3)
            self.text(
                s, self.watermark, x, self.Inches(0.1), w, self.Inches(0.35), 11, color="999999"
            )
        return s

    def text(
        self,
        s,
        text: str,
        x,
        y,
        w,
        h,
        size: float,
        bold=False,
        color=None,
        bullets=None,
        align=None,
    ):
        from pptx.dml.color import RGBColor

        box = s.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        lines = bullets if bullets is not None else text.split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            self.runs(p, ("• " if bullets is not None else "") + line, size, bold, color)
            if align is not None:
                p.alignment = align
            if bullets is not None:
                p.space_after = self.Pt(6)
        if color is not None:
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = RGBColor.from_string(color)
        return box

    def runs(self, p, text: str, size: float, bold: bool, color) -> None:
        """Markdown **bold** and *italic* as runs; other markup as plain text."""
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
            if not part:
                continue
            r = p.add_run()
            if part.startswith("**"):
                r.text, r.font.bold = part[2:-2], True
            elif part.startswith("*"):
                r.text, r.font.italic = part[1:-1], True
            elif part.startswith("`"):
                r.text = part[1:-1]
                r.font.name = "Consolas"
            else:
                r.text = part
            r.font.size = self.Pt(size)
            if self.neutral and not part.startswith("`"):
                r.font.name = "Arial"  # present on every system; the default theme font is not
            if bold:
                r.font.bold = True

    def picture(self, s, png: bytes, x, y, w, h) -> None:
        """The image at its native aspect ratio, centered in the box."""
        from ..importers.content import png_size

        iw, ih = png_size(png)
        scale = min(w / iw, h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        s.shapes.add_picture(io.BytesIO(png), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)

    def footnote(self, s, text: str) -> None:
        self.text(
            s,
            text,
            self.margin,
            self.H - self.Inches(0.55),
            self.W - 2 * self.margin,
            self.Inches(0.4),
            11,
            color="666666",
        )

    # -- slides -----------------------------------------------------------------------

    def title_slide(self) -> None:
        r = self.report
        s = self.slide(r.name)
        lines = [r.description] if r.description else []
        details = " · ".join(
            f"{label}: {v}"
            for label, v in (
                ("Author", r.author),
                ("Date", r.date),
                ("Data cut-off", r.data_cutoff),
            )
            if v
        )
        if details:
            lines.append(details)
        if lines:
            self.text(
                s,
                "\n".join(lines),
                self.margin,
                self.top + self.Inches(0.3),
                self.W - 2 * self.margin,
                self.Inches(2),
                18,
                color="444444",
            )

    def room(self, room, notes: str) -> None:
        s = self.slide(room.name, notes)
        if room.description:
            self.text(
                s,
                room.description,
                self.margin,
                self.top + self.Inches(0.3),
                self.W - 2 * self.margin,
                self.Inches(2),
                20,
                color="444444",
            )

    def payload(self, v: Visualization) -> dict[str, Any] | None:
        if self.levels.get(v.id, "data") == "preview":
            return None  # reduced to its image in the export review
        return self.report._payloads.get(v.id)

    def kpi_tile(self, s, v: Visualization, x, y, w, h) -> None:
        from pptx.enum.text import PP_ALIGN

        p = self.payload(v)
        if not p:
            return
        tile = min(h, self.Inches(2.6))  # a compact tile, centered in the cell
        y, h = y + (h - tile) // 2, tile
        value = f"{p['display']} {p['unit']}".strip()
        self.text(s, v.name, x, y, w, h * 0.25, 16, color="444444", align=PP_ALIGN.CENTER)
        self.text(s, value, x, y + h * 0.25, w, h * 0.45, 40, bold=True, align=PP_ALIGN.CENTER)
        if p.get("delta") is not None:
            good = (p["delta"] > 0) == p["higher_is_better"]
            arrow = "▲" if p["delta"] > 0 else "▼" if p["delta"] < 0 else "■"
            self.text(s, f"{arrow} {p['delta_display']}", x, y + h * 0.7, w, h * 0.25, 18,
                      color="1A7F37" if good else "C62828", align=PP_ALIGN.CENTER)  # fmt: skip

    def table(self, s, v: Visualization, x, y, w, h) -> None:
        p = self.payload(v)
        rows, cols = p["rows"][: MAX_TABLE[0]], p["columns"][: MAX_TABLE[1]]
        shape = s.shapes.add_table(
            len(rows) + 1, len(cols), x, y, w, min(h, self.Inches(0.37) * (len(rows) + 1))
        )
        t = shape.table
        for j, c in enumerate(cols):
            t.cell(0, j).text = str(c)
        for i, r in enumerate(rows, 1):
            for j in range(len(cols)):
                val = r[j]
                t.cell(i, j).text = (
                    "" if val is None else f"{val:,.6g}" if isinstance(val, float) else str(val)
                )
        for cell in (t.cell(i, j) for i in range(len(rows) + 1) for j in range(len(cols))):
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = self.Pt(12)
                    if self.neutral:
                        run.font.name = "Arial"
        if p["total_rows"] > len(rows) or len(p["columns"]) > len(cols):
            self.footnote(
                s,
                f"{len(rows)} of {p['total_rows']:,} rows and {len(cols)} of "
                f"{len(p['columns'])} columns. The full table is in the DataWalk report.",
            )

    def text_panel(self, s, v: Visualization, x, y, w, h) -> None:
        p = self.payload(v)
        lines = []
        for raw in p["markdown"].splitlines():
            line = re.sub(r"^\s*#{1,6}\s*", "", raw).strip()
            line = re.sub(r"^\s*([-*+]|\d+[.)])\s+", "• ", line)
            line = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", line)
            if line:
                lines.append(line.rstrip("\\").rstrip())
        evidence = [
            e.name for k in p.get("links", []) if (e := self.report.get(k)) and e.status == "active"
        ]
        if evidence:
            lines.append(f"*Evidence: {', '.join(evidence)}*")
        box = self.text(s, "", x, y, w, h, 24)
        tf = box.text_frame
        tf.paragraphs[0].text = ""
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            self.runs(para, line, 24 if len(lines) < 6 else 18, False, None)
            para.space_after = self.Pt(8)

    def visualization(self, v: Visualization, notes: str) -> None:
        s = self.slide(v.name, notes)
        top, bottom = self.top, self.H - self.Inches(0.7)
        findings = v.overrides.explanation.findings
        kind = v.source.chart_type
        full_w = self.W - 2 * self.margin
        side = self.Inches(3.9) if findings else 0
        box_w = full_w - side - (self.Inches(0.3) if side else 0)
        if kind == "kpi" and self.payload(v):
            self.kpi_tile(s, v, self.margin, top + self.Inches(0.6), box_w, self.Inches(3.4))
        elif kind == "table" and self.payload(v):
            self.table(s, v, self.margin, top, box_w, bottom - top)
        elif kind == "text" and self.payload(v):
            self.text_panel(s, v, self.margin, top, full_w, bottom - top)
        else:
            png = self.images.get(v.id) or self.report._previews.get(v.id)
            if png:
                self.picture(s, png, self.margin, top, box_w, bottom - top)
            else:
                self.text(
                    s,
                    "No image of this chart is available.",
                    self.margin,
                    top,
                    box_w,
                    self.Inches(1),
                    18,
                    color="666666",
                )
        if findings:
            x = self.margin + box_w + self.Inches(0.3)
            self.text(s, "Key findings", x, top, side, self.Inches(0.5), 18, bold=True)
            self.text(
                s,
                "",
                x,
                top + self.Inches(0.55),
                side,
                bottom - top - self.Inches(0.55),
                16,
                bullets=list(findings),
            )
        if v.source.mode == "interactive":
            self.footnote(s, INTERACTIVE_NOTE)

    def dashboard(self, d: Dashboard, notes: str) -> None:
        s = self.slide(d.name, notes)
        members = [
            v
            for i in d.items
            if (v := next((x for x in self.report.visualizations if x.id == i), None))
            and v.status == "active"
        ]
        n = len(members)
        cols = n if n <= 3 else int(-(-(n**0.5) // 1))
        rows = -(-n // cols)
        top, bottom = self.top, self.H - self.Inches(0.7)
        gap = self.Inches(0.25)
        cw = (self.W - 2 * self.margin - gap * (cols - 1)) // cols
        ch = (bottom - top - gap * (rows - 1)) // rows
        for i, v in enumerate(members):
            x = self.margin + (i % cols) * (cw + gap)
            y = top + (i // cols) * (ch + gap)
            if v.source.chart_type == "kpi" and self.payload(v):
                self.kpi_tile(s, v, x, y, cw, ch)
            elif png := self.images.get(v.id) or self.report._previews.get(v.id):
                self.picture(s, png, x, y, cw, ch)
        if any(v.source.mode == "interactive" for v in members):
            self.footnote(s, INTERACTIVE_NOTE)


def needs_images(report: Report) -> list[str]:
    """Charts in the tour that have no stored image: the viewer renders them (headless)."""
    ids = []
    for _, _, target in report.tour_stops():
        members = target.items if isinstance(target, Dashboard) else [target.id] if target else []
        for vid in members:
            v = next(x for x in report.visualizations if x.id == vid)
            if v.id not in report._previews and v.source.chart_type not in ("kpi", "table", "text"):
                ids.append(v.id)
    return list(dict.fromkeys(ids))


def write_pptx(
    report: Report,
    path: Path,
    images: dict[str, bytes],
    template: Path | None = None,
    levels: dict[str, str] | None = None,
    watermark: str = "",
) -> Path:
    try:
        import pptx  # noqa: F401
    except ImportError as e:
        raise ImportError("PPTX export needs python-pptx: pip install 'datawalk[pptx]'") from e
    deck = _Deck(report, template, images, levels or {}, watermark)
    deck.title_slide()
    for stop, room, target in report.tour_stops():
        if target is None:
            deck.room(room, stop.notes)
        elif isinstance(target, Dashboard):
            deck.dashboard(target, stop.notes)
        else:
            deck.visualization(target, stop.notes)
    tmp = Path(path).with_suffix(".tmp.pptx")
    deck.prs.save(str(tmp))
    shutil.move(tmp, path)
    return Path(path)
